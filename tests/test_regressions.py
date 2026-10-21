"""回归测试：离线页面、API 状态和真实 Chromium；不访问亚马逊。"""
import asyncio
import io
import json
import re
import tempfile
import threading
import time
import unittest
from contextlib import ExitStack
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import Error as PlaywrightError, sync_playwright

import server
from scraper import browser, crawler, login
from scraper.errors import CrawlError
from scraper.parser import _rating_from_text, parse_page_reviews
from marketplaces import reviews_url, MARKETPLACES, STAR_FILTERS
from task_manager import TaskManager
from results import combination_key, unique_reviews
from scraper.product import extract_product, safe_image
from scraper.origin import review_origin

TASK = {"asin": "B0G92M6VFM", "country": "us", "mode": "deep", "cap": 10}


def review(rid, body):
    return f'<div data-hook="review" id="{rid}"><a data-hook="review-title"><span>Title {rid}</span></a><i data-hook="review-star-rating" class="a-star-4"><span class="a-icon-alt">4.5 out of 5 stars</span></i><span data-hook="review-body"><span>{body}</span><span>second paragraph</span></span></div>'


def dated_review(rid, date):
    return review(rid, 'origin fixture').replace('</div>', '<span data-hook="review-date">'+date+'</span></div>')

class ReviewOriginTests(unittest.TestCase):
    def test_english_origin_names_for_all_markets(self):
        names = {"us":"the United States","uk":"the United Kingdom","de":"Germany","fr":"France","it":"Italy","es":"Spain","jp":"Japan","ca":"Canada","mx":"Mexico","br":"Brazil","au":"Australia","in":"India","nl":"the Netherlands","se":"Sweden","pl":"Poland","sg":"Singapore","ae":"the United Arab Emirates","tr":"Turkey","sa":"Saudi Arabia","be":"Belgium"}
        self.assertEqual(set(names),set(MARKETPLACES))
        for code,name in names.items():
            with self.subTest(country=code):
                self.assertEqual(review_origin(f"Reviewed in {name} on 7 April 2026"),(code,name))

    def test_localized_origin_and_unknown_regions(self):
        cases = [("Reviewed\u00a0in\u00a0Australia\u00a0on 7 April 2026","au"),("Rezension aus Deutschland vom 7. April 2026","de"),("Commenté au Royaume-Uni le 7 avril 2026","uk"),("Revisado en España el 7 de abril de 2026","es"),("Recensito in Italia il 7 aprile 2026","it"),("Avaliado no Brasil em 7 de abril de 2026","br"),("Beoordeeld in Nederland op 7 april 2026","nl"),("Recenserad i Sverige den 7 april 2026","se"),("Opinia napisana w Polsce dnia 7 kwietnia 2026","pl"),("日本で2026年4月7日にレビュー済み","jp"),("2026年4月7日在澳大利亚评论","au"),("تمت المراجعة في المملكة العربية السعودية في 7 أبريل 2026","sa"),("7 Nisan 2026 tarihinde Türkiye’de değerlendirildi","tr")]
        for date,code in cases:
            with self.subTest(date=date):
                self.assertEqual(review_origin(date)[0],code)
        for date in ("",None,"7 April 2026","Reviewed in Indiana on 7 April 2026","Reviewed in Atlantis on 7 April 2026","Australia"):
            self.assertEqual(review_origin(date)[0],"")
        self.assertEqual(review_origin("Reviewed in Atlantis on 7 April 2026")[1],"Atlantis")

class StateAndAPITests(unittest.TestCase):
    def setUp(self):
        self.manager = TaskManager()
        self.manager_patch = patch.object(server, "manager", self.manager)
        self.manager_patch.start()
        self.client = TestClient(server.app)

    def tearDown(self):
        self.client.close()
        self.manager_patch.stop()

    def test_invalid_input_rejected_without_worker(self):
        for body in [
            {"asins": "bad"}, {"asins": ["BAD"]},
            {"asins": [TASK["asin"]], "country": "../us"},
            {"asins": [TASK["asin"]], "mode": "unknown"},
            {"asins": [TASK["asin"]], "cap": -1},
            {"asins": [TASK["asin"]], "cap": 1.2},
            {"asins": [TASK["asin"]], "star_key": "bad"},
            {"asins": [TASK["asin"]], "local_only": "false"},
            {"asins": [TASK["asin"]], "local_only": 1},
        ]:
            with self.subTest(body=body):
                response = self.client.post("/api/start", json=body)
                self.assertEqual(response.status_code, 422)
                self.assertFalse(response.json()["ok"])
                self.assertFalse(self.manager.running)

    def test_failure_is_not_completed_and_export_records_reason(self):
        with patch.object(server, "crawl", side_effect=CrawlError("blocked", "verification required")):
            response = self.client.post("/api/start", json={"asins": [TASK["asin"]]})
            self.assertEqual(response.status_code, 200)
            self.manager.worker.result(timeout=5)
        state = self.client.get("/api/status").json()
        self.assertFalse(state["running"])
        self.assertEqual(state["phase"], "failed")
        self.assertEqual((state["done"], state["total"]), (1, 1))
        self.assertEqual(state["outcomes"][TASK["asin"]]["status"], "blocked")
        workbook = load_workbook(io.BytesIO(self.client.get("/api/export").content))
        rows = list(workbook["总览"].values)
        self.assertEqual(dict(zip(rows[0], rows[1]))["任务状态"], "blocked")
        self.assertEqual(dict(zip(rows[0], rows[1]))["状态说明"], "verification required")


    def test_batch_closes_browser_once_before_completion_for_all_endings(self):
        tasks = [{**TASK, "id":"first"}, {**TASK, "id":"second"}]
        for ending in ("success", "empty", "blocked", "unexpected", "cancelled"):
            with self.subTest(ending=ending):
                manager = TaskManager()
                manager.start_run(tasks)
                sample = [{"body":"preserved result","asin":TASK["asin"]}]
                def collect(task, **kwargs):
                    close.assert_not_called()
                    self.assertTrue(kwargs["keep_browser"])
                    owners.append(threading.get_ident())
                    if ending == "blocked":
                        raise CrawlError("blocked", "verification required", sample)
                    if ending == "unexpected":
                        raise RuntimeError("unexpected fixture error")
                    if ending == "cancelled":
                        manager.stop()
                    return [] if ending == "empty" else sample
                def cleanup():
                    owners.append(threading.get_ident())
                    self.assertTrue(manager.running)
                owners = []
                with patch.object(server, "crawl", side_effect=collect), patch.object(browser, "close", side_effect=cleanup) as close, patch.object(server.logger, "exception"):
                    server.BROWSER_WORKER.submit(server._worker_run, tasks, manager).result(timeout=5)
                close.assert_called_once()
                self.assertEqual(len(set(owners)), 1)
                self.assertFalse(manager.running)
                self.assertEqual(manager.phase, {"blocked":"partial","unexpected":"failed","cancelled":"cancelled"}.get(ending,"completed"))
                self.assertEqual(manager.done, 1 if ending=="cancelled" else 2)
                if ending not in ("empty", "unexpected"):
                    self.assertEqual(manager.results["first"], sample)

    def test_browser_cleanup_error_keeps_task_results_and_completion(self):
        manager = TaskManager()
        manager.start_run([TASK])
        with patch.object(server, "crawl", return_value=[{"body":"preserved"}]), patch.object(browser, "close", side_effect=RuntimeError("already disconnected")), patch.object(server.logger, "exception") as log:
            server.BROWSER_WORKER.submit(server._worker_run, [TASK], manager).result(timeout=5)
        self.assertFalse(manager.running)
        self.assertEqual(manager.phase, "completed")
        self.assertEqual(manager.results[TASK["asin"]][0]["body"], "preserved")
        log.assert_called_once()


    def test_local_country_option_per_market_and_filtered_empty_reason(self):
        observed = []
        def collect(task,**kwargs):
            observed.append(task)
            return []
        with patch.object(browser,"is_logged_in",return_value=True),patch.object(server,"crawl",side_effect=collect):
            response = self.client.post("/api/start",json={"asins":[TASK["asin"]],"countries":["uk","us"],"mode":"deep","local_only":True})
            self.manager.worker.result(timeout=5)
        self.assertEqual(response.status_code,200)
        self.assertEqual([t["country"] for t in observed],["uk","us"])
        self.assertTrue(all(t["local_only"] for t in observed))
        self.assertEqual(self.manager.status,"任务结束：没有符合条件的评论")
        self.assertTrue(all("符合本站地区条件" in o["message"] for o in self.manager.outcomes.values()))
        with patch.object(server,"crawl",return_value=[]):
            self.client.post("/api/start",json={"asins":[TASK["asin"]]})
            self.manager.worker.result(timeout=5)
        self.assertFalse(self.manager.tasks[0]["local_only"])

    def test_review_origin_export_and_legacy_results_do_not_confuse_market_with_origin(self):
        date = "Reviewed in Australia on 7 April 2026"
        task = {**TASK,"country":"uk","local_only":False}
        raw = {"review_id":"AU1","body":"from Australia","date":date}
        merged = unique_reviews([task],{TASK["asin"]:[raw]})
        self.assertEqual((merged[0]["country"],merged[0]["review_country"]),("uk","au"))
        self.assertNotIn("review_country",raw)
        workbook = load_workbook(io.BytesIO(server.build_excel([task],{TASK["asin"]:[raw]})))
        overview = list(workbook["总览"].values)
        self.assertEqual(dict(zip(*overview))["评价来源"],"全部地区")
        rows = list(workbook[TASK["asin"]].values)
        exported = dict(zip(*rows))
        self.assertEqual(exported["国家/地区"],"英国")
        self.assertEqual(exported["评价标注地区"],"澳大利亚")
        self.assertEqual(exported["日期"],date)

    def test_partial_reviews_preserved(self):
        sample = {"asin": TASK["asin"], "country": "us", "body": "first page"}
        with patch.object(server, "crawl", side_effect=CrawlError("pagination_error", "timeout", [sample])):
            self.client.post("/api/start", json={"asins": [TASK["asin"]]})
            self.manager.worker.result(timeout=5)
        self.assertEqual(self.manager.phase, "partial")
        self.assertEqual(self.manager.results[TASK["asin"]], [sample])

    def test_empty_and_success_and_consecutive_runs(self):
        for reviews in [[], [{"asin": TASK["asin"], "rating": 4.0, "body": "text"}]]:
            with patch.object(server, "crawl", return_value=reviews):
                response = self.client.post("/api/start", json={"asins": [TASK["asin"].lower(), TASK["asin"]]})
                self.manager.worker.result(timeout=5)
            self.assertEqual(response.json()["total"], 1)
            self.assertEqual(self.manager.phase, "completed")
            self.assertEqual(self.manager.done, 1)
            self.assertEqual(self.manager.outcomes[TASK["asin"]]["status"], "success" if reviews else "empty")

    def test_login_and_run_exclusive_and_second_save_enabled(self):
        self.manager.start_login("us")
        self.assertEqual(self.client.post("/api/start", json={"asins": [TASK["asin"]]}).status_code, 409)
        self.assertEqual(self.client.post("/api/login", json={"country": "us"}).status_code, 409)
        self.assertEqual(self.client.post("/api/login/save").status_code, 409)
        self.manager.on_login_ready("us")
        self.assertEqual(self.client.post("/api/login/save").status_code, 200)
        self.manager.on_login_saved("us")
        self.manager.start_login("us")
        self.manager.on_login_ready("us")
        self.assertEqual(self.client.post("/api/login/save").status_code, 200)
        self.client.post("/api/login/cancel")
        self.assertTrue(self.manager.login_cancel.is_set())

    def test_start_is_atomic_across_threads(self):
        barrier = threading.Barrier(2)
        admitted = []
        def start():
            barrier.wait()
            try:
                self.manager.start_run([TASK])
                admitted.append(True)
            except ValueError:
                admitted.append(False)
        threads = [threading.Thread(target=start) for _ in range(2)]
        for thread in threads: thread.start()
        for thread in threads: thread.join()
        self.assertEqual(sorted(admitted), [False, True])

    def test_cancel_marks_unprocessed_and_preserves_count(self):
        self.manager.start_run([TASK, {**TASK, "asin": "B07XJ8C8F5"}])
        self.manager.add_result(TASK["asin"], [{"body": "kept"}])
        self.manager.stop()
        self.manager.mark_complete()
        self.assertEqual(self.manager.phase, "cancelled")
        self.assertEqual((self.manager.done, self.manager.total), (1, 2))
        self.assertEqual(self.manager.outcomes["B07XJ8C8F5"]["status"], "cancelled")

    def test_sse_reconnect_starts_with_current_snapshot(self):
        async def run():
            self.manager.start_run([TASK])
            self.manager.add_result(TASK["asin"], [{"body": "kept"}])
            self.manager.mark_complete()
            for _ in range(2):
                queue = asyncio.Queue()
                self.manager.subscribe(queue, asyncio.get_running_loop())
                event = await queue.get()
                self.assertEqual(event["type"], "snapshot")
                self.assertEqual(event["done"], 1)
                self.assertEqual(event["results"][TASK["asin"]][0]["body"], "kept")
                self.manager.unsubscribe(queue)
        asyncio.run(run())

    def test_server_login_and_crawl_share_one_worker_thread(self):
        owners = []
        def fake_login(*args, **kwargs):
            owners.append(threading.get_ident())
            self.assertTrue(kwargs["keep_browser"])
            self.assertEqual(kwargs["verification_asin"], TASK["asin"])
        def fake_crawl(*args, **kwargs):
            owners.append(threading.get_ident())
            self.assertTrue(kwargs["keep_browser"])
            return [{"body": "read in same session"}]
        with patch.object(server.login_mod, "run_login", side_effect=fake_login):
            response = self.client.post("/api/login", json={"country": "us", "asin": TASK["asin"]})
            self.assertEqual(response.status_code, 200)
            server.BROWSER_WORKER.submit(lambda: None).result(timeout=5)
        with patch.object(server, "crawl", side_effect=fake_crawl):
            self.client.post("/api/start", json={"asins": [TASK["asin"]]})
            self.manager.worker.result(timeout=5)
        self.assertEqual(self.manager.login_phase, "saved")
        self.assertEqual(len(owners), 2)
        self.assertEqual(owners[0], owners[1])

    def test_invalid_login_target_rejected(self):
        response = self.client.post("/api/login", json={"country": "us", "asin": "wrong"})
        self.assertEqual(response.status_code, 422)

    def test_multiselect_cartesian_product_and_deduplicated_export(self):
        asins = [TASK["asin"], "B07XJ8C8F5"]
        def collect(task, **kwargs):
            return [{"review_id": "R1", "rating": 4, "body": "same review across filters"}]
        with patch.object(browser, "is_logged_in", return_value=True), patch.object(server, "crawl", side_effect=collect):
            response = self.client.post("/api/start", json={"asins": asins, "countries": ["us", "uk"], "star_keys": ["4 星", "好评（4-5 星）"], "mode": "deep", "cap": 10})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["total"], 8)
            self.manager.worker.result(timeout=5)
        state = self.client.get("/api/status").json()
        self.assertEqual(len(state["results"]), 8)
        self.assertEqual(len(state["reviews"]), 4)
        self.assertEqual(len(state["outcomes"]), 8)
        self.assertEqual(state["done"], 8)
        self.assertEqual(len(state["reviews"][0]["matched_filters"]), 2)
        workbook = load_workbook(io.BytesIO(self.client.get("/api/export").content))
        self.assertEqual(workbook["总览"].max_row, 9)
        self.assertEqual(set(workbook.sheetnames), {"总览", *[f"{asin}_{country}" for asin in asins for country in ["us", "uk"]]})
        self.assertTrue(all(sheet.max_row == 2 for sheet in workbook.worksheets[1:]))

    def test_multiselect_missing_session_lists_exact_market(self):
        with patch.object(browser, "is_logged_in", side_effect=lambda code: code == "us"):
            response = self.client.post("/api/start", json={"asins": [TASK["asin"]], "countries": ["us", "uk"], "star_keys": ["1 星"], "mode": "deep"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["missing_countries"], ["uk"])

    def test_invalid_multi_selections_and_task_limit(self):
        for fields in [{"countries": []}, {"countries": ["bad"]}, {"star_keys": []},
                       {"star_keys": ["全部", "1 星"]}, {"star_keys": ["unknown"]},
                       {"countries": list(MARKETPLACES), "star_keys": ["1 星", "2 星"], "asins": [f"B{i:09d}" for i in range(20)]}]:
            with self.subTest(fields=fields):
                response = self.client.post("/api/start", json={"asins": [TASK["asin"]], "mode": "deep", **fields})
                self.assertEqual(response.status_code, 422)

    def test_quick_multi_market_does_not_create_star_combinations(self):
        with patch.object(server, "crawl", return_value=[]):
            response = self.client.post("/api/start", json={"asins": [TASK["asin"]], "countries": ["us", "us", "uk"], "star_keys": ["1 星", "2 星"], "mode": "quick"})
            self.manager.worker.result(timeout=5)
        self.assertEqual(response.json()["total"], 2)
        self.assertEqual({task["star_key"] for task in self.manager.tasks}, {"全部"})

    def test_result_cache_restores_completed_task_without_login_state(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "results.json"
            first = TaskManager(path)
            first.start_run([TASK])
            first.add_result(TASK["asin"], [{"body": "preserved", "rating": 5}])
            first.update_product(TASK["asin"], "us", {"title": "Saved product"})
            first.mark_complete()
            second = TaskManager(path)
            self.assertFalse(second.running)
            self.assertIsNone(second.login_phase)
            self.assertEqual(second.snapshot()["reviews"][0]["body"], "preserved")
            self.assertEqual(second.products[TASK["asin"]+":us"]["title"], "Saved product")


class SessionTests(unittest.TestCase):
    def test_reviews_url_has_canonical_slash_and_star_query(self):
        self.assertEqual(reviews_url(TASK["asin"], "us"), f'https://www.amazon.com/product-reviews/{TASK["asin"]}/')
        self.assertEqual(reviews_url(TASK["asin"], "us", "1 星"), f'https://www.amazon.com/product-reviews/{TASK["asin"]}/?filterByStar=one_star')

    def test_corrupt_expired_and_verified_state(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(browser, "STORAGE_DIR", Path(folder)):
            path = Path(folder) / "login_state_us.json"
            self.assertEqual(browser.session_status("us"), "missing")
            path.write_text("bad", encoding="utf-8")
            self.assertEqual(browser.session_status("us"), "invalid")
            path.write_text(json.dumps({"cookies": [{"name": "at-main", "value": "test", "expires": time.time()-10}]}))
            self.assertEqual(browser.session_status("us"), "expired")
            path.write_text(json.dumps({"cookies": [{"name": "at-main", "value": "test", "expires": time.time()+1000}]}))
            self.assertEqual(browser.session_status("us"), "saved")
            browser.mark_session("us", "expired")
            self.assertFalse(browser.is_logged_in("us"))
            browser.mark_session("us", "verified")
            self.assertEqual(browser.session_status("us"), "verified")

    def test_localized_rating(self):
        for text, expected in [("4.5 out of 5 stars", 4.5), ("4,5 von 5 Sternen", 4.5),
                               ("5つ星のうち4.0", 4), ("", 0)]:
            self.assertEqual(_rating_from_text(text), expected)


class BrowserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pw = sync_playwright().start()
        cls.engine = cls.pw.chromium.launch(headless=True)

    @classmethod
    def tearDownClass(cls):
        cls.engine.close()
        cls.pw.stop()

    def setUp(self):
        self.context = self.engine.new_context()
        self.page = self.context.new_page()
        self.page.set_default_timeout(3000)

    def tearDown(self):
        self.context.close()

    def load_frontend(self, origin=False):
        html = Path("static/index.html").read_text(encoding="utf-8")
        css = Path("static/app.css").read_text(encoding="utf-8")
        script = Path("static/app.js").read_text(encoding="utf-8").replace("\ninit();", "\n// disable network init")
        translations = Path("static/i18n.js").read_text(encoding="utf-8")
        html = re.sub(r'<link rel="stylesheet" href="/static/app\.css[^\"]*">', lambda _: '<style>'+css+'</style>', html)
        html = re.sub(r'<script src="/static/i18n\.js[^\"]*"></script>', lambda _: '<script>'+translations+'</script>', html)
        html = re.sub(r'<script src="/static/app\.js[^\"]*"></script>', lambda _: '<script>'+script+'</script>', html)
        if origin:
            self.page.route('http://ui.test/**', lambda route: route.fulfill(content_type='text/html', body=html))
            self.page.goto('http://ui.test/')
        else:
            self.page.set_content(html)
        self.page.evaluate("c => {config=c; connected=true;renderCountryOptions();renderStarOptions();renderEstimate();}", {"marketplaces": list(MARKETPLACES.values()), "star_filters": list(STAR_FILTERS)})

    def test_product_metadata_uses_real_image_and_rejects_external_url(self):
        self.page.set_content('<h1 id="productTitle">Real product title</h1><img id="landingImage" src="https://m.media-amazon.com/images/I/example.jpg"><a id="bylineInfo">Example Brand</a>')
        item = extract_product(self.page, TASK["asin"], "us")
        self.assertEqual(item["title"], "Real product title")
        self.assertEqual(item["image"], "https://m.media-amazon.com/images/I/example.jpg")
        self.assertEqual(safe_image("https://evil.example/pixel"), "")
        self.assertEqual(safe_image("javascript:alert(1)"), "")

    def test_frontend_checkbox_multi_selection_payload(self):
        self.load_frontend()
        self.page.locator('#asins').fill(TASK["asin"]+"\nB07XJ8C8F5")
        self.page.locator('[data-mode="deep"]').click()
        self.page.locator('input[name="countries"][value="uk"]').check()
        self.page.locator('input[name="countries"][value="de"]').check()
        self.page.locator('input[name="star_keys"][value="1 星"]').check()
        self.page.locator('input[name="star_keys"][value="2 星"]').check()
        self.assertFalse(self.page.locator('input[name="star_keys"][value="全部"]').is_checked())
        self.assertEqual(self.page.locator('#estimatedCount').inner_text(), "12 个任务")
        self.page.evaluate("""() => {window.fetch=async (url, options) => {
          if(url==='/api/start') window.sentBody=JSON.parse(options.body);
          return {ok:true,json:async()=>url==='/api/start'?{ok:true}:{...state,revision:99,instance_id:'test'}};
        };}""")
        self.page.locator('#startBtn').click()
        self.page.wait_for_function('window.sentBody !== undefined')
        payload = self.page.evaluate('window.sentBody')
        self.assertEqual(payload['countries'], ['us','uk','de'])
        self.assertEqual(payload['star_keys'], ['1 星','2 星'])

    def test_frontend_review_search_pagination_and_sort(self):
        self.load_frontend()
        manager = TaskManager()
        manager.start_run([TASK])
        manager.add_result(TASK['asin'], [{'review_id':str(i),'title':f'Review {i}', 'body':'needle' if i==18 else 'ordinary text','author':'Tester','rating':(i%5)+1} for i in range(30)])
        manager.mark_complete()
        self.page.evaluate('s=>applySnapshot(s)', manager.snapshot())
        self.assertEqual(self.page.locator('#resultBody .review-row').count(), 12)
        self.page.locator('#pagination [data-page="2"]').first.click()
        self.assertIn('13–24 / 30', self.page.locator('#pagination').inner_text())
        self.page.locator('#reviewSearch').fill('needle')
        self.assertEqual(self.page.locator('#resultBody .review-row').count(), 1)
        self.assertIn('Review 18', self.page.locator('#resultBody').inner_text())
        self.page.locator('#reviewSearch').fill('')
        self.page.locator('#reviewSort').select_option('lowest')
        self.assertIn('1.0 / 5', self.page.locator('.review-rating').first.inner_text())


    def test_language_switch_preserves_filters_and_review_reading_state(self):
        self.load_frontend()
        self.page.locator('#asins').fill(TASK['asin'])
        self.page.locator('[data-mode="deep"]').click()
        self.page.locator('input[name="countries"][value="uk"]').check()
        self.page.locator('input[name="star_keys"][value="2 星"]').check()
        manager = TaskManager()
        manager.start_run([TASK])
        manager.add_result(TASK['asin'], [{'review_id':str(i), 'title':'原始评论标题', 'body':'source review '*25, 'rating':2} for i in range(30)])
        manager.mark_complete()
        self.page.evaluate('s=>applySnapshot(s)', manager.snapshot())
        self.page.locator('#reviewSearch').fill('source')
        self.page.locator('#resultStar').select_option('2')
        self.page.locator('#reviewSort').select_option('lowest')
        self.page.locator('#pagination [data-page="2"]').first.click()
        self.page.locator('.expand-review').first.click()
        self.page.locator('#languageSelect').select_option('en')
        self.assertEqual(self.page.locator('html').get_attribute('lang'), 'en')
        self.assertEqual(self.page.locator('#asins').input_value(), TASK['asin'])
        self.assertEqual(self.page.locator('#reviewSearch').input_value(), 'source')
        self.assertEqual(self.page.locator('#resultStar').input_value(), '2')
        self.assertEqual(self.page.locator('#reviewSort').input_value(), 'lowest')
        self.assertIn('13–24 / 30', self.page.locator('#pagination').inner_text())
        self.assertEqual(self.page.locator('.expand-review').first.inner_text().strip(), 'Show less')
        self.assertEqual(self.page.locator('.review-title').first.inner_text(), '原始评论标题')
        self.assertEqual(self.page.evaluate('[...selectedCountries]'), ['us','uk'])
        self.assertEqual(self.page.evaluate('[...selectedStars]'), ['2 星'])
        self.assertIn('2 tasks', self.page.locator('#estimatedCount').inner_text())
        self.page.locator('#languageSelect').select_option('zh')
        self.assertIn('收起内容', self.page.locator('.expand-review').first.inner_text())
        self.assertEqual(self.page.locator('h1').inner_text(), '亚马逊商品工具箱')

    def test_english_protocol_payload_and_backend_messages(self):
        self.load_frontend()
        self.page.locator('#asins').fill(TASK['asin'])
        self.page.locator('[data-mode="deep"]').click()
        self.page.locator('input[name="star_keys"][value="1 星"]').check()
        self.page.locator('#languageSelect').select_option('en')
        self.page.evaluate("""() => {window.fetch=async(url,options)=>{
          if(url==='/api/start') window.sentBody=JSON.parse(options.body);
          return {ok:true,json:async()=>url==='/api/start'?{ok:true}:{...state,revision:99,instance_id:'test'}};
        }}""")
        self.page.locator('#startBtn').click()
        self.page.wait_for_function('window.sentBody !== undefined')
        self.assertEqual(self.page.evaluate('window.sentBody.star_keys'), ['1 星'])
        self.assertEqual(self.page.evaluate('window.sentBody.countries'), ['us'])
        translated = self.page.evaluate("""() => [
            localizeServer('ASIN B0G92M6VFM | 第 2 页 | 已抓 24 条 | 第 2 批新增 12 条'),
            localizeServer('请先完成这些站点的登录：美国、英国'),
            localizeServer('评论列表页要求再次认证；请点左侧登录，在评论页完成认证后再次保存','review_auth_required')
        ]""")
        self.assertIn('Batch 2', translated[0])
        self.assertIn('24 reviews', translated[0])
        self.assertIn('United States, United Kingdom', translated[1])
        self.assertIn('Market sessions', translated[2])
        self.assertFalse(any(re.search(r'[\u3400-\u9fff]', item) for item in translated))

    def test_top_navigation_and_white_readable_responsive_layout(self):
        self.load_frontend()
        self.page.locator('#toggleSetup').click()
        self.assertTrue(self.page.locator('#setupContent').is_hidden())
        self.page.locator('#navSetup').click()
        self.assertTrue(self.page.locator('#setupContent').is_visible())
        self.page.locator('#navSessions').click()
        self.assertTrue(self.page.locator('#sessionsPanel').evaluate('e=>e.open'))
        self.assertEqual(self.page.locator('.sidebar').count(), 0)
        self.assertEqual(self.page.locator('body').evaluate('e=>getComputedStyle(e).backgroundColor'), 'rgb(255, 255, 255)')
        self.assertGreaterEqual(float(self.page.locator('body').evaluate('e=>parseFloat(getComputedStyle(e).fontSize)')),16)
        self.page.locator('#languageSelect').select_option('en')
        for width in (1440, 1100, 800, 420, 360):
            with self.subTest(width=width):
                self.page.set_viewport_size({'width':width,'height':950})
                self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'))
                self.assertTrue(self.page.locator('#languageSelect').is_visible())
        ratios = self.page.evaluate("""() => {
          const luminance = hex => {
            const c=hex.match(/[a-f0-9]{2}/gi).map(v=>parseInt(v,16)/255).map(v=>v<=.04045?v/12.92:((v+.055)/1.055)**2.4);
            return c[0]*.2126+c[1]*.7152+c[2]*.0722;
          };
          return ['#172033','#263247','#526176','#2556e8'].map(c=>(1.05)/(luminance(c)+.05));
        }""")
        self.assertTrue(all(ratio>=4.5 for ratio in ratios))

    def test_language_preference_survives_reload(self):
        self.load_frontend(origin=True)
        self.page.locator('#languageSelect').select_option('en')
        self.page.reload()
        self.assertEqual(self.page.locator('html').get_attribute('lang'), 'en')
        self.assertEqual(self.page.evaluate("t('{count} 个 ASIN',{count:1})"), '1 ASIN')
        self.assertEqual(self.page.locator('#languageSelect').input_value(), 'en')
        self.assertEqual(self.page.locator('h1').inner_text(), 'Amazon Product Toolkit')
        self.page.evaluate("setLanguage('zh')")
        self.page.reload()
        self.assertEqual(self.page.locator('html').get_attribute('lang'), 'zh-CN')

    def test_all_static_labels_have_english_translations(self):
        self.load_frontend()
        missing = self.page.evaluate("""() => [...document.querySelectorAll('[data-i18n],[data-i18n-placeholder],[data-i18n-aria],[data-i18n-title]')].flatMap(e=>Array.from(e.attributes).filter(a=>a.name.startsWith('data-i18n')).map(a=>a.value)).filter(key=>!EN[key])""")
        self.assertEqual(missing, [])
        self.page.locator('#languageSelect').select_option('en')
        self.page.locator('#toggleMarkets').click()
        self.assertEqual(self.page.locator('#countryOptions input').count(), len(MARKETPLACES))
        self.assertFalse(re.search(r'[\u3400-\u9fff]', self.page.locator('#setupPanel').inner_text()))

    def crawl_fixture(self, html, task=None, handler=None, cancel_event=None):
        self.page.route("https://www.amazon.*/**", handler or (
            lambda route: route.fulfill(status=200, content_type="text/html", body=html)))
        with ExitStack() as stack:
            for name, value in [("launch", self.page), ("is_logged_in", True),
                                ("load_storage_state", self.page), ("close", None),
                                ("mark_session", None)]:
                stack.enter_context(patch.object(browser, name, return_value=value))
            return crawler.crawl(task or TASK, cancel_event=cancel_event or threading.Event())

    def test_classic_navigation_collects_second_page(self):
        first = review("R1", "first") + '<li class="a-last"><a href="/second">Next</a></li>'
        second = review("R2", "second")
        def handler(route):
            html = second if route.request.url.endswith("/second") else first
            route.fulfill(status=200, content_type="text/html", body=html)
        result = self.crawl_fixture(first, handler=handler)
        self.assertEqual([r["review_id"] for r in result], ["R1", "R2"])
        self.assertIn("second paragraph", result[0]["body"])

    def test_append_pagination_deduplicates(self):
        second = json.dumps(review("R2", "second"))
        html = review("R1", "first") + f'<button data-hook="show-more-button" onclick=\'document.getElementById("list").insertAdjacentHTML("beforeend", {second});this.remove()\'>More</button>'
        result = self.crawl_fixture('<div id="list">' + html + '</div>')
        self.assertEqual([r["review_id"] for r in result], ["R1", "R2"])


    def test_local_origin_continues_foreign_page_and_counts_cap_after_filter(self):
        first = dated_review("AU1","Reviewed in Australia on 7 April 2026")+'<li class="a-last"><a href="/second">Next</a></li>'
        second = dated_review("UK1","Reviewed in the United Kingdom on 8 April 2026")+dated_review("UK2","Reviewed in the United Kingdom on 9 April 2026")
        result = self.crawl_fixture(first,{**TASK,"country":"uk","cap":1,"local_only":True},handler=lambda route:route.fulfill(content_type="text/html",body=second if route.request.url.endswith("/second") else first))
        self.assertEqual([r["review_id"] for r in result],["UK1"])
        self.assertEqual((result[0]["review_country"],result[0]["country"]),("uk","uk"))

    def test_local_origin_append_pagination_and_foreign_duplicates(self):
        second = json.dumps(dated_review("UK1","Reviewed in the United Kingdom on 8 April 2026"))
        html = dated_review("AU1","Reviewed in Australia on 7 April 2026")+f'<button data-hook="show-more-button" onclick=\'document.getElementById("list").insertAdjacentHTML("beforeend", {second});this.remove()\'>More</button>'
        result = self.crawl_fixture('<div id="list">'+html+'</div>',{**TASK,"country":"uk","local_only":True})
        self.assertEqual([r["review_id"] for r in result],["UK1"])

    def test_local_origin_excludes_unknown_and_can_be_disabled(self):
        html = dated_review("AU1","Reviewed in Australia on 7 April 2026")+dated_review("UK1","Reviewed in the United Kingdom on 8 April 2026")+dated_review("UNKNOWN","")
        for enabled,ids in ((True,["UK1"]),(False,["AU1","UK1","UNKNOWN"])):
            self.page.unroute_all()
            with self.subTest(enabled=enabled):
                result = self.crawl_fixture(html,{**TASK,"country":"uk","local_only":enabled})
                self.assertEqual([r["review_id"] for r in result],ids)

    def test_local_origin_all_foreign_and_repeated_page_guard(self):
        html = dated_review("AU1","Reviewed in Australia on 7 April 2026")
        self.assertEqual(self.crawl_fixture(html,{**TASK,"country":"uk","local_only":True}),[])
        self.page.unroute_all()
        first = html+'<li class="a-last"><a href="/second">Next</a></li>'
        second = html+'<li class="a-last"><a href="/third">Next</a></li>'
        with self.assertRaises(CrawlError) as caught:
            self.crawl_fixture(first,{**TASK,"country":"uk","local_only":True},handler=lambda route:route.fulfill(content_type="text/html",body=first if "product-reviews" in route.request.url else second))
        self.assertEqual(caught.exception.status,"pagination_error")
        self.assertEqual(caught.exception.reviews,[])

    def test_local_origin_failure_keeps_only_matching_partial_reviews(self):
        first = dated_review("AU1","Reviewed in Australia on 7 April 2026")+dated_review("UK1","Reviewed in the United Kingdom on 8 April 2026")+'<li class="a-last"><a href="/second">Next</a></li>'
        def handler(route):
            route.fulfill(status=503,body="restricted") if route.request.url.endswith("/second") else route.fulfill(content_type="text/html",body=first)
        with self.assertRaises(CrawlError) as caught:
            self.crawl_fixture(first,{**TASK,"country":"uk","local_only":True},handler=handler)
        self.assertEqual([r["review_id"] for r in caught.exception.reviews],["UK1"])

    def test_origin_setting_and_review_country_display_bilingual(self):
        self.load_frontend(origin=True)
        self.page.locator('#asins').fill(TASK["asin"])
        self.page.locator('[data-mode="deep"]').click()
        self.assertFalse(self.page.locator('#localOnly').is_checked())
        self.page.locator('#localOnly').check()
        self.page.wait_for_function("JSON.parse(localStorage.getItem('reviewStudio.settings')||'{}').localOnly===true")
        self.page.locator('#languageSelect').select_option('en')
        self.assertTrue(self.page.locator('#localOnly').is_checked())
        self.page.evaluate("""() => {window.fetch=async(url,options)=>{
          if(url==='/api/start') window.sentBody=JSON.parse(options.body);
          return {ok:true,json:async()=>url==='/api/start'?{ok:true}:{...state,revision:99,instance_id:'test'}};
        }}""")
        self.page.locator('#startBtn').click()
        self.page.wait_for_function('window.sentBody !== undefined')
        self.assertTrue(self.page.evaluate('window.sentBody.local_only'))
        manager = TaskManager()
        manager.start_run([{**TASK,"country":"uk"}])
        manager.add_result(TASK["asin"],[{"body":"AU review","date":"Reviewed in Australia on 7 April 2026"}])
        manager.mark_complete()
        self.page.evaluate('s=>applySnapshot({...s,instance_id:"new"})',manager.snapshot())
        self.assertIn('Review country: Australia',self.page.locator('.origin-pill').inner_text())
        self.assertIn('United Kingdom market',self.page.locator('.review-meta .market-pill').inner_text())
        self.assertEqual(self.page.locator('.review-date').get_attribute('title'),'Reviewed in Australia on 7 April 2026')
        self.page.locator('#languageSelect').select_option('zh')
        self.assertIn('评价地区：澳大利亚',self.page.locator('.origin-pill').inner_text())

    def test_cap_honored(self):
        result = self.crawl_fixture(review("R1", "first") + review("R2", "second"), {**TASK, "cap": 1})
        self.assertEqual(len(result), 1)

    def test_quick_mode_ignores_deep_mode_cap(self):
        result = self.crawl_fixture(review("R1", "first") + review("R2", "second"),
                                    {**TASK, "mode": "quick", "cap": 1})
        self.assertEqual(len(result), 2)

    def test_http_errors_are_not_empty_success(self):
        for http_status, expected in [(404, "http_error"), (429, "blocked")]:
            with self.subTest(http_status=http_status):
                self.page.unroute_all()
                with self.assertRaises(CrawlError) as caught:
                    self.crawl_fixture("", handler=lambda route: route.fulfill(status=http_status, body="error"))
                self.assertEqual(caught.exception.status, expected)

    def test_anonymous_save_does_not_overwrite_existing_state(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(browser, "STORAGE_DIR", Path(folder)), patch.object(browser, "current_context", return_value=self.page.context):
            path = Path(folder) / "login_state_us.json"
            path.write_text("existing saved state", encoding="utf-8")
            with self.assertRaises(RuntimeError):
                browser.save_storage_state("us")
            self.assertEqual(path.read_text(encoding="utf-8"), "existing saved state")

    def test_quick_captcha_is_explicit_failure(self):
        with self.assertRaises(CrawlError) as caught:
            self.crawl_fixture('<form action="/errors/validateCaptcha"><input id="captchacharacters"></form>',
                               {**TASK, "mode": "quick"})
        self.assertEqual(caught.exception.status, "blocked")

    def test_login_redirect_invalidates_session(self):
        with self.assertRaises(CrawlError) as caught:
            self.crawl_fixture('<form name="signIn"><input id="ap_email"></form>')
        self.assertEqual(caught.exception.status, "review_auth_required")

    def test_unknown_page_not_treated_as_empty(self):
        with patch.object(crawler, "_wait_reviews", return_value=False):
            with self.assertRaises(CrawlError) as caught:
                self.crawl_fixture('<p>unexpected markup</p>')
        self.assertEqual(caught.exception.status, "parse_error")

    def test_explicit_empty_page(self):
        with patch.object(crawler, "_wait_reviews", return_value=False):
            result = self.crawl_fixture('<div data-hook="no-reviews-message">No customer reviews</div>')
        self.assertEqual(result, [])

    def test_pagination_timeout_keeps_partial_data(self):
        with patch.object(crawler, "_wait_page_change", return_value=False):
            with self.assertRaises(CrawlError) as caught:
                self.crawl_fixture(review("R1", "first") + '<button data-hook="show-more-button">More</button>')
        self.assertEqual(caught.exception.status, "pagination_error")
        self.assertEqual(len(caught.exception.reviews), 1)

    def test_cancel_during_pagination_keeps_partial_data(self):
        event = threading.Event()
        def cancel(*args, **kwargs):
            event.set()
            raise CrawlError("cancelled", "stopped")
        with patch.object(crawler, "_wait_page_change", side_effect=cancel):
            with self.assertRaises(CrawlError) as caught:
                self.crawl_fixture(review("R1", "first") + '<button data-hook="show-more-button">More</button>',
                                   cancel_event=event)
        self.assertEqual(caught.exception.status, "cancelled")
        self.assertEqual(len(caught.exception.reviews), 1)

    def test_missing_rating_does_not_wait_for_absent_element(self):
        self.page.set_content('<div data-hook="review"><span data-hook="review-body">body</span></div>')
        start = time.monotonic()
        parsed = parse_page_reviews(self.page, TASK["asin"], "us", 1)
        self.assertEqual(parsed[0]["rating"], 0)
        self.assertLess(time.monotonic() - start, 2)

    def test_login_cancel_and_timeout_release_browser(self):
        self.page.route("https://www.amazon.com/**", lambda route: route.fulfill(body='<p>home</p>'))
        for cancelled in [True, False]:
            ready, done, cancel = threading.Event(), threading.Event(), threading.Event()
            if cancelled: cancel.set()
            with patch.object(browser, "launch", return_value=self.page), patch.object(browser, "is_logged_in", return_value=False), patch.object(browser, "close") as close:
                with self.assertRaises(RuntimeError):
                    login.run_login("us", ready, done, cancel_event=cancel, timeout=0)
                close.assert_called_once()

    def test_login_verifies_in_separate_tab_despite_frontend_redirect(self):
        def respond(route):
            body = ('<script>location.replace("/?ref_=nav_ya_signin")</script>'
                    if "/gp/css/" in route.request.url else '<a href="/gp/flex/sign-out.html">Sign out</a>')
            route.fulfill(status=200, content_type="text/html", body=body)
        self.page.context.route("https://www.amazon.com/**", respond)
        self.page.goto("https://www.amazon.com/")
        original_url = self.page.url
        login._verify_login(self.page, "us", threading.Event())
        self.assertEqual(self.page.url, original_url)
        self.assertFalse(self.page.is_closed())
        self.assertEqual(len(self.page.context.pages), 1)

    def test_exact_interrupted_goto_error_checks_final_page(self):
        self.page.route("https://www.amazon.com/**", lambda route: route.fulfill(
            status=200, content_type="text/html", body="<p>Ready</p>"))
        goto = self.page.goto
        def interrupt(*args, **kwargs):
            goto("https://www.amazon.com/?ref_=nav_ya_signin", wait_until="domcontentloaded")
            raise PlaywrightError('Page.goto: Navigation to "https://www.amazon.com/gp/css/homepage.html" is interrupted by another navigation to "https://www.amazon.com/?ref_=nav_ya_signin"')
        with patch.object(self.page, "goto", side_effect=interrupt):
            response = login._navigate(self.page, "https://www.amazon.com/gp/css/homepage.html", threading.Event())
        self.assertEqual(response.status, 200)
        self.assertEqual(self.page.url, "https://www.amazon.com/?ref_=nav_ya_signin")

    def test_public_account_home_is_not_login_verification(self):
        self.context.route("https://www.amazon.com/**", lambda route: route.fulfill(
            status=200, content_type="text/html", body='<h1>Your Account</h1><a data-nav-role="signin">Sign in</a>'))
        with self.assertRaises(login.LoginChallenge):
            login._verify_login(self.page, "us", threading.Event())

    def test_verification_reads_target_product_reviews(self):
        paths = []
        def respond(route):
            paths.append(route.request.url)
            route.fulfill(status=200, content_type="text/html", body=review("R1", "real review fixture"))
        self.context.route("https://www.amazon.com/**", respond)
        login._verify_login(self.page, "us", threading.Event(), TASK["asin"])
        self.assertTrue(any(f'/product-reviews/{TASK["asin"]}' in path for path in paths))
        self.assertFalse(any('/gp/css/' in path for path in paths))

    def test_login_opens_requested_reviews_before_waiting_for_user(self):
        destinations = []
        def navigate(page, url, event):
            destinations.append(url)
        with patch.object(browser, "launch", return_value=self.page), patch.object(browser, "is_logged_in", return_value=False), patch.object(browser, "close"), patch.object(login, "_navigate", side_effect=navigate):
            with self.assertRaises(RuntimeError):
                login.run_login("us", threading.Event(), threading.Event(), timeout=0,
                                verification_asin=TASK["asin"])
        self.assertEqual(destinations, [reviews_url(TASK["asin"], "us")])

    def test_review_challenge_is_preserved_for_manual_completion(self):
        def respond(route):
            if "/product-reviews/" in route.request.url:
                route.fulfill(status=302, headers={"Location": "https://www.amazon.com/ap/signin"})
            else:
                route.fulfill(status=200, content_type="text/html", body='<form name="signIn"><input id="ap_email"></form>')
        self.context.route("https://www.amazon.com/**", respond)
        with self.assertRaises(login.LoginChallenge) as caught:
            login._verify_login(self.page, "us", threading.Event(), TASK["asin"])
        self.assertIn("评论列表页", str(caught.exception))
        self.assertFalse(caught.exception.page.is_closed())
        self.assertIn("/ap/signin", caught.exception.page.url)

    def test_login_verification_still_rejects_signin_redirect(self):
        def respond(route):
            if "/gp/css/" in route.request.url:
                route.fulfill(status=302, headers={"Location": "https://www.amazon.com/ap/signin"})
            else:
                route.fulfill(status=200, content_type="text/html", body='<form name="signIn"><input id="ap_email"></form>')
        self.page.context.route("https://www.amazon.com/**", respond)
        with self.assertRaises(CrawlError) as caught:
            login._verify_login(self.page, "us", threading.Event())
        self.assertEqual(caught.exception.status, "review_auth_required")
        self.assertFalse(self.page.is_closed())
        self.assertEqual(len(self.page.context.pages), 2)

    def test_verification_retry_keeps_window_and_can_save(self):
        self.page.route("https://www.amazon.com/**", lambda route: route.fulfill(
            status=200, content_type="text/html", body="<p>Login window</p>"))
        ready, done = threading.Event(), threading.Event()
        messages = []
        def retry(message):
            self.assertFalse(self.page.is_closed())
            messages.append(message)
            done.set()
        with patch.object(browser, "launch", return_value=self.page), patch.object(browser, "is_logged_in", return_value=False), patch.object(browser, "close") as close, patch.object(browser, "save_storage_state") as save, patch.object(login, "_verify_login", side_effect=[RuntimeError("Please finish login"), None]):
            login.run_login("us", ready, done, on_ready=done.set, on_retry=retry)
            save.assert_called_once_with("us")
            close.assert_called_once()
        self.assertEqual(messages, ["Please finish login"])

    def test_retry_can_be_cancelled_without_saving(self):
        self.page.route("https://www.amazon.com/**", lambda route: route.fulfill(
            status=200, content_type="text/html", body="<p>Login window</p>"))
        ready, done, cancel = threading.Event(), threading.Event(), threading.Event()
        def retry(message):
            cancel.set()
            done.set()
        with patch.object(browser, "launch", return_value=self.page), patch.object(browser, "is_logged_in", return_value=False), patch.object(browser, "close"), patch.object(browser, "save_storage_state") as save, patch.object(login, "_verify_login", side_effect=RuntimeError("not ready")):
            with self.assertRaisesRegex(RuntimeError, "登录已取消"):
                login.run_login("us", ready, done, on_ready=done.set, on_retry=retry, cancel_event=cancel)
            save.assert_not_called()

    def test_retry_adopts_challenge_page_and_retains_saved_browser(self):
        self.page.route("https://www.amazon.com/**", lambda route: route.fulfill(
            status=200, content_type="text/html", body="<p>Login window</p>"))
        challenge = self.context.new_page()
        ready, done = threading.Event(), threading.Event()
        failure = login.LoginChallenge("Complete review authentication", challenge)
        with patch.object(browser, "launch", return_value=self.page), patch.object(browser, "is_logged_in", return_value=False), patch.object(browser, "close") as close, patch.object(browser, "set_current_page") as adopt, patch.object(browser, "save_storage_state") as save, patch.object(login, "_verify_login", side_effect=[failure, None]) as verify:
            login.run_login("us", ready, done, on_ready=done.set,
                            on_retry=lambda message: done.set(), keep_browser=True)
            adopt.assert_called_once_with(challenge)
            self.assertIs(verify.call_args_list[1].args[0], challenge)
            save.assert_called_once_with("us")
            close.assert_not_called()
        self.assertFalse(challenge.is_closed())

    def test_network_error_is_not_treated_as_successful_redirect(self):
        with patch.object(self.page, "goto", side_effect=PlaywrightError("net::ERR_NAME_NOT_RESOLVED")):
            with self.assertRaises(PlaywrightError):
                login._navigate(self.page, "https://www.amazon.com/", threading.Event())

    def test_frontend_restores_results_and_handles_repeated_login(self):
        self.load_frontend()
        state = TaskManager()
        state.start_run([TASK])
        state.add_result(TASK["asin"], [{"asin": TASK["asin"], "body": "<script>unsafe</script>", "rating": 4}])
        state.mark_complete()
        self.page.evaluate("s => applySnapshot(s)", {**state.snapshot(), "sessions": {"us": "saved"}})
        self.assertEqual(self.page.locator("#progressBar").evaluate("e=>e.style.width"), "100%")
        self.assertIn("1 / 1", self.page.locator("#countLabel").inner_text())
        self.assertIn("待验证", self.page.locator("#loginBadge").text_content())
        self.assertEqual(self.page.locator("#resultBody .review-row").count(), 1)
        self.assertEqual(self.page.locator("#resultBody script").count(), 0)
        for _ in range(2):
            state.start_login("us")
            state.on_login_ready("us")
            self.page.evaluate("s => applySnapshot(s)", state.snapshot())
            self.assertFalse(self.page.locator("#loginSaveBtn").is_disabled())
            state.finish_login()
            self.page.evaluate("s => applySnapshot(s)", state.snapshot())
            self.assertTrue(self.page.locator("#loginSaveBtn").is_disabled())
            state.on_login_retry("us", "登录页面仍在跳转")
            self.page.evaluate("s => applySnapshot(s)", state.snapshot())
            self.assertFalse(self.page.locator("#loginSaveBtn").is_disabled())
            self.assertIn("认证页面已保留", self.page.locator("#loginBannerText").inner_text())
            state.finish_login()
            state.on_login_saved("us")
        snapshot = state.snapshot()
        snapshot["revision"] = 0
        snapshot["done"] = 0
        self.page.evaluate("s => applySnapshot(s)", snapshot)
        self.assertIn("1 / 1", self.page.locator("#countLabel").inner_text())


class ThreadIsolationTests(unittest.TestCase):

    def test_batch_reuses_live_markets_then_closes_all_windows_and_keeps_session_files(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(browser, "STORAGE_DIR", Path(folder)), patch.dict("os.environ", {"AMAZON_REVIEW_HEADLESS":"1"}):
            executor = ThreadPoolExecutor(max_workers=1)
            def run():
                pages, contexts, saved = {}, {}, {}
                tasks = [{**TASK, "country":code, "id":code} for code in ("us","uk")]
                for code in ("us","uk"):
                    pages[code] = browser.launch(country=code)
                    contexts[code] = browser.current_context()
                    contexts[code].route("https://**.amazon.*/**", lambda route: route.fulfill(content_type="text/html", body=review("R1","batch fixture")))
                    contexts[code].add_cookies([{"name":"at-main","value":"offline-fixture","url":"https://"+MARKETPLACES[code]["domain"]+"/"}])
                    browser.save_storage_state(code)
                    saved[code] = (Path(folder)/f"login_state_{code}.json").read_bytes()
                original = crawler.crawl
                def collect(task, **kwargs):
                    self.assertIs(browser.launch(country=task["country"]), pages[task["country"]])
                    self.assertIs(browser.current_context(), contexts[task["country"]])
                    return original(task, **kwargs)
                manager = TaskManager()
                manager.start_run(tasks)
                with patch.object(server, "crawl", side_effect=collect):
                    server._worker_run(tasks, manager)
                self.assertFalse(browser.is_running())
                self.assertIsNone(browser.current_context())
                self.assertTrue(all(page.is_closed() for page in pages.values()))
                self.assertEqual(manager.phase, "completed")
                self.assertEqual(len(manager.snapshot()["reviews"]), 2)
                for code in saved:
                    self.assertEqual((Path(folder)/f"login_state_{code}.json").read_bytes(), saved[code])
                # 关闭后下一批仍可从已保存的会话创建新浏览器。
                page = browser.launch(country="us")
                self.assertFalse(page.is_closed())
                self.assertTrue(any(c["name"]=="at-main" for c in browser.current_context().cookies()))
            try:
                executor.submit(run).result(timeout=25)
            finally:
                executor.submit(browser.close).result(timeout=10)
                executor.shutdown()

    def test_product_metadata_closes_own_browser_but_preserves_live_login_browser(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(browser, "STORAGE_DIR", Path(folder)), patch.dict("os.environ", {"AMAZON_REVIEW_HEADLESS":"1"}):
            executor = ThreadPoolExecutor(max_workers=1)
            def run():
                with patch.object(crawler, "_goto", side_effect=lambda page,*args,**kwargs:page.set_content('<h1 id="productTitle">Fixture product</h1>')):
                    info = crawler.fetch_product_info(TASK["asin"],"us")
                    self.assertEqual(info["title"],"Fixture product")
                    self.assertFalse(browser.is_running())
                    original = browser.launch(country="us")
                    context = browser.current_context()
                    crawler.fetch_product_info(TASK["asin"],"us")
                    self.assertTrue(browser.is_running())
                    self.assertIs(browser.current_context(),context)
                    self.assertFalse(original.is_closed())
                    self.assertEqual(len(context.pages),1)
                browser.close()
                with patch.object(crawler, "_goto", side_effect=RuntimeError("metadata failure")):
                    with self.assertRaisesRegex(RuntimeError,"metadata failure"):
                        crawler.fetch_product_info(TASK["asin"],"us")
                self.assertFalse(browser.is_running())
            try:
                executor.submit(run).result(timeout=25)
            finally:
                executor.submit(browser.close).result(timeout=10)
                executor.shutdown()

    def test_switching_markets_preserves_each_live_context(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(browser, "STORAGE_DIR", Path(folder)):
            try:
                first = browser.launch(headless=True, country="us")
                us_context = browser.current_context()
                second = browser.launch(headless=True, country="uk")
                uk_context = browser.current_context()
                self.assertIsNot(us_context, uk_context)
                self.assertFalse(first.is_closed())
                self.assertFalse(second.is_closed())
                self.assertIs(browser.launch(country="us"), first)
                self.assertIs(browser.current_context(), us_context)
                self.assertIs(browser.launch(country="uk"), second)
            finally:
                browser.close()

    def test_login_then_crawl_reuses_live_context_without_cookie_replay(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(browser, "STORAGE_DIR", Path(folder)), patch.dict("os.environ", {"AMAZON_REVIEW_HEADLESS": "1"}):
            executor = ThreadPoolExecutor(max_workers=1)
            def sign_in():
                page = browser.launch(headless=True, country="us")
                context = browser.current_context()
                context.route("https://www.amazon.com/**", lambda route: route.fulfill(
                    status=200, content_type="text/html", body=review("R1", "same live context")))
                context.add_cookies([{"name": "at-main", "value": "offline-fixture", "url": "https://www.amazon.com/"}])
                done = threading.Event()
                login.run_login("us", threading.Event(), done, on_ready=done.set,
                                verification_asin=TASK["asin"], keep_browser=True)
                return id(context)
            def scrape():
                with patch.object(browser._local.browser, "new_context", side_effect=AssertionError("Live login context must not be replaced")):
                    result = crawler.crawl({**TASK, "cap": 1}, keep_browser=True)
                return id(browser.current_context()), result
            try:
                original = executor.submit(sign_in).result(timeout=15)
                current, result = executor.submit(scrape).result(timeout=15)
                self.assertEqual(current, original)
                self.assertEqual(len(result), 1)
                self.assertIn("same live context", result[0]["body"])
            finally:
                executor.submit(browser.close).result(timeout=10)
                executor.shutdown()

    def test_independent_browser_objects_in_two_threads(self):
        barrier = threading.Barrier(2)
        results, errors = [], []
        def worker(title):
            try:
                page = browser.launch(headless=True)
                page.set_content(f"<title>{title}</title>")
                barrier.wait(timeout=15)
                results.append(page.title())
            except Exception as exc:
                errors.append(str(exc))
                barrier.abort()
            finally:
                browser.close()
        threads = [threading.Thread(target=worker, args=(title,)) for title in ["first", "second"]]
        for thread in threads: thread.start()
        for thread in threads: thread.join(timeout=25)
        self.assertFalse(any(thread.is_alive() for thread in threads))
        self.assertEqual(errors, [])
        self.assertEqual(sorted(results), ["first", "second"])


if __name__ == "__main__":
    unittest.main()
