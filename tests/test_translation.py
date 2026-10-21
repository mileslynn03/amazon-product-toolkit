"""翻译回归使用本地模拟 HTTP 服务和 Chromium，不调用付费模型。"""
import copy
import io
import json
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from playwright.sync_api import sync_playwright

import server
import translation
import test_regressions as old_tests
from task_manager import TaskManager
from translation import TranslationSettings, TranslationError, TranslationCancelled, translate_batch, translate_results

TASK = {"id":"first","asin":"B0G92M6VFM","country":"us","mode":"quick","star_key":"全部","cap":None}
SECRET = "offline-fixture-key-not-real"


class MockProvider:
    def __enter__(self):
        self.calls, self.mode = [], "normal"
        owner = self
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args):
                pass
            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                owner.calls.append({"path":self.path,"authorization":self.headers.get("Authorization"),"x_api_key":self.headers.get("x-api-key"),"anthropic_version":self.headers.get("anthropic-version"),"body":body})
                if owner.mode == "unauthorized":
                    self.send_response(401); self.end_headers()
                    self.wfile.write(json.dumps({"error":SECRET}).encode())
                    return
                if owner.mode == "redirect":
                    self.send_response(302); self.send_header("Location","/leaked"); self.end_headers()
                    return
                native = self.path.endswith("/messages")
                items = json.loads(body["messages"][0 if native else 1]["content"])["reviews"]
                rows = [{"id":r["id"],"title":"Translated "+r["title"] if r["title"] else "","body":"Translated "+r["body"] if r["body"] else ""} for r in reversed(items)]
                if owner.mode == "invalid":
                    rows = [{"id":"wrong","title":"wrong","body":"wrong"}]
                if owner.mode == "partial" and any("bad" in r["body"] for r in items):
                    rows[0]["body"] = ""
                content = json.dumps({"translations":rows},ensure_ascii=False)
                if owner.mode == "fenced":
                    content = "```json\n"+content+"\n```"
                if native:
                    blocks = [{"type":"text","text":content}]
                    if owner.mode == "thinking":
                        split = len(content)//2
                        blocks = [{"type":"thinking","thinking":"private fixture thinking"},{"type":"text","text":content[:split]},{"type":"text","text":content[split:]}]
                    stop = owner.mode if owner.mode in {"max_tokens","refusal","pause_turn","tool_use","model_context_window_exceeded"} else "end_turn"
                    response = {"type":"message","role":"assistant","content":blocks,"stop_reason":stop}
                else:
                    response = {"choices":[{"finish_reason":"stop","message":{"content":content}}]}
                data = json.dumps(response,ensure_ascii=False).encode()
                self.send_response(200); self.send_header("Content-Type","application/json"); self.end_headers(); self.wfile.write(data)
        self.http = ThreadingHTTPServer(("127.0.0.1",0),Handler)
        self.thread = threading.Thread(target=self.http.serve_forever,daemon=True)
        self.thread.start()
        self.base = f"http://127.0.0.1:{self.http.server_port}/v1"
        return self

    def __exit__(self,*args):
        self.http.shutdown(); self.http.server_close(); self.thread.join()

    def settings(self,**kwargs):
        return TranslationSettings(base_url=self.base,model="fixture-model",api_key=SECRET,**kwargs)

    def body(self,**kwargs):
        return {"base_url":self.base,"model":"fixture-model","api_key":SECRET,**kwargs}


def seed(manager,body="good"):
    manager.start_run([TASK])
    manager.add_result("first",[{"review_id":"R1","title":"Original title","body":body,"rating":5}])
    manager.mark_complete()


class TranslationTransportTests(unittest.TestCase):
    def test_endpoint_validation_and_secret_repr(self):
        with MockProvider() as provider:
            settings = provider.settings()
            self.assertEqual(settings.endpoint,provider.base+"/chat/completions")
            self.assertNotIn(SECRET,repr(settings))
            full = TranslationSettings(**provider.body(base_url=provider.base+"/chat/completions/"))
            self.assertEqual(full.endpoint,settings.endpoint)
            for value in ("file:///tmp/key","https://user:pass@example.com/v1","https://example.com/v1?key=secret","https://example.com/#token"):
                with self.subTest(value=value),self.assertRaises(ValidationError):
                    TranslationSettings(**provider.body(base_url=value))
            for field,value in (("batch_size",1.2),("timeout","90"),("target_language","unknown")):
                with self.assertRaises(ValidationError):
                    TranslationSettings(**provider.body(**{field:value}))

    def test_http_payload_auth_and_shuffled_ids_preserve_matching(self):
        with MockProvider() as provider:
            items = [{"id":"0","title":"One","body":"first"},{"id":"1","title":"Two","body":"second"}]
            output = translate_batch(provider.settings(),items,threading.Event())
            self.assertEqual(output["0"]["body"],"Translated first")
            self.assertEqual(output["1"]["title"],"Translated Two")
            self.assertEqual(provider.calls[0]["authorization"],"Bearer "+SECRET)
            self.assertEqual(provider.calls[0]["path"],"/v1/chat/completions")
            self.assertNotIn(SECRET,json.dumps(provider.calls[0]["body"]))
            self.assertNotIn("temperature",provider.calls[0]["body"])
            self.assertFalse(provider.calls[0]["body"]["stream"])

    def test_code_fences_and_invalid_output(self):
        with MockProvider() as provider:
            items = [{"id":"0","title":"Title","body":"Body"}]
            provider.mode = "fenced"
            self.assertEqual(translate_batch(provider.settings(),items,threading.Event())["0"]["body"],"Translated Body")
            provider.mode = "invalid"
            with self.assertRaises(TranslationError):
                translate_batch(provider.settings(),items,threading.Event())

    def test_auth_errors_redacted_and_redirect_never_forwards_credentials(self):
        with MockProvider() as provider:
            for mode in ("unauthorized","redirect"):
                provider.mode = mode
                with self.assertRaises(TranslationError) as caught:
                    translate_batch(provider.settings(),[{"id":"0","title":"","body":"text"}],threading.Event())
                self.assertTrue(caught.exception.fatal)
                self.assertNotIn(SECRET,str(caught.exception))
            self.assertEqual(len(provider.calls),2)
            self.assertFalse(any(call["path"]=="/leaked" for call in provider.calls))

    def test_transient_retries_and_cancellation(self):
        settings = TranslationSettings(base_url="http://127.0.0.1/v1",model="fixture")
        event = Mock(spec=threading.Event)
        event.is_set.return_value=False
        event.wait.return_value=False
        with patch.object(translation,"_request",side_effect=[TranslationError("temporary",retryable=True),TranslationError("temporary",retryable=True),{"0":{"title":"","body":"ok"}}]) as request:
            output = translate_batch(settings,[{"id":"0","title":"","body":"text"}],event)
        self.assertEqual(output["0"]["body"],"ok")
        self.assertEqual(request.call_count,3)
        event.wait.return_value=True
        with patch.object(translation,"_request",side_effect=TranslationError("temporary",retryable=True)):
            with self.assertRaises(TranslationCancelled):
                translate_batch(settings,[],event)

    def test_truncated_output_and_missing_fields_are_not_success(self):
        settings = TranslationSettings(base_url="http://127.0.0.1/v1",model="fixture")
        for choice in ({"finish_reason":"length","message":{"content":"{}"}},{"message":{"content":json.dumps({"translations":[{"id":"0","title":"","body":""}]})}}):
            response = Mock()
            response.read.return_value=json.dumps({"choices":[choice]}).encode()
            response.__enter__=Mock(return_value=response);response.__exit__=Mock(return_value=False)
            opener=Mock();opener.open.return_value=response
            with patch.object(translation,"build_opener",return_value=opener),self.assertRaises(TranslationError):
                translate_batch(settings,[{"id":"0","title":"","body":"must keep"}],threading.Event())


class AnthropicTests(unittest.TestCase):
    def test_native_endpoint_forms_and_wrong_protocol_validation(self):
        for base,expected in (
            ("https://api.anthropic.com","https://api.anthropic.com/v1/messages"),
            ("https://api.anthropic.com/v1/","https://api.anthropic.com/v1/messages"),
            ("https://gateway.example/anthropic","https://gateway.example/anthropic/v1/messages"),
            ("https://gateway.example/api/messages/","https://gateway.example/api/messages"),
        ):
            settings=TranslationSettings(provider="anthropic",base_url=base,model="fixture")
            self.assertEqual(settings.endpoint,expected)
        for fields in ({"provider":"invalid"},{"provider":"anthropic","base_url":"https://example.com/v1/chat/completions"},{"base_url":"https://example.com/v1/messages"},{"max_tokens":127},{"max_tokens":"4096"}):
            with self.subTest(fields=fields),self.assertRaises(ValidationError):
                TranslationSettings(**{"base_url":"https://example.com/v1","model":"fixture",**fields})

    def test_native_headers_system_max_tokens_and_shuffled_review_mapping(self):
        with MockProvider() as provider:
            settings=provider.settings(provider="anthropic",max_tokens=8192)
            items=[{"id":"0","title":"First","body":"first review"},{"id":"1","title":"Second","body":"second review"}]
            output=translate_batch(settings,items,threading.Event())
            call=provider.calls[0]
            self.assertEqual(call["path"],"/v1/messages")
            self.assertEqual(call["x_api_key"],SECRET)
            self.assertEqual(call["anthropic_version"],"2023-06-01")
            self.assertIsNone(call["authorization"])
            self.assertEqual(call["body"]["max_tokens"],8192)
            self.assertIsInstance(call["body"]["system"],str)
            self.assertEqual([m["role"] for m in call["body"]["messages"]],["user"])
            self.assertEqual(output["0"]["body"],"Translated first review")
            self.assertEqual(output["1"]["title"],"Translated Second")
            self.assertNotIn(SECRET,json.dumps(call["body"]))

    def test_thinking_blocks_are_ignored_and_text_blocks_joined(self):
        with MockProvider() as provider:
            provider.mode="thinking"
            result=translate_batch(provider.settings(provider="anthropic"),[{"id":"0","title":"","body":"review"}],threading.Event())
            self.assertEqual(result["0"]["body"],"Translated review")
            self.assertNotIn("thinking",json.dumps(result))

    def test_incomplete_stop_reasons_and_malformed_native_response_are_failures(self):
        with MockProvider() as provider:
            for mode in ("max_tokens","model_context_window_exceeded","refusal","pause_turn","tool_use","invalid"):
                provider.mode=mode
                with self.subTest(mode=mode),self.assertRaises(TranslationError):
                    translate_batch(provider.settings(provider="anthropic"),[{"id":"0","title":"","body":"review"}],threading.Event())

    def test_native_auth_error_and_redirect_do_not_leak_key(self):
        with MockProvider() as provider:
            for mode in ("unauthorized","redirect"):
                provider.mode=mode
                with self.assertRaises(TranslationError) as error:
                    translate_batch(provider.settings(provider="anthropic"),[{"id":"0","title":"","body":"review"}],threading.Event())
                self.assertNotIn(SECRET,str(error.exception))
                self.assertTrue(error.exception.fatal)
            self.assertEqual(len(provider.calls),2)
            self.assertTrue(all(call["authorization"] is None for call in provider.calls))

    def test_legacy_openai_cache_preserved_and_native_cache_is_separate(self):
        import hashlib
        with MockProvider() as provider:
            review={"title":"Title","body":"Body"}
            settings=provider.settings()
            legacy=hashlib.sha256(json.dumps([settings.endpoint,settings.model,settings.target_language,review["title"],review["body"]],ensure_ascii=False).encode()).hexdigest()
            self.assertEqual(settings.cache_key(review),legacy)
            self.assertNotEqual(settings.cache_key(review),provider.settings(provider="anthropic").cache_key(review))
            manager=TaskManager();seed(manager)
            native=provider.settings(provider="anthropic")
            manager.start_translation(native);translate_results(manager,native)
            self.assertEqual(manager.translation["phase"],"completed")
            self.assertEqual(manager.translation["provider"],"anthropic")
            self.assertEqual(manager.results["first"][0]["translation"]["provider"],"anthropic")
            self.assertEqual(manager.results["first"][0]["body"],"good")
            count=len(provider.calls)
            manager.start_translation(native);translate_results(manager,native)
            self.assertEqual(len(provider.calls),count)
            rows=list(load_workbook(io.BytesIO(server.build_excel(manager.tasks,manager.results)))[TASK["asin"]].values)
            self.assertEqual(dict(zip(rows[0],rows[1]))["翻译接口协议"],"Anthropic Messages")

    def test_native_test_and_translation_endpoints(self):
        manager=TaskManager();seed(manager)
        with MockProvider() as provider,patch.object(server,"manager",manager):
            client=TestClient(server.app)
            try:
                result=client.post("/api/llm/test",json=provider.body(provider="anthropic"))
                self.assertEqual(result.status_code,200)
                result=client.post("/api/translate",json=provider.body(provider="anthropic",max_tokens=8192))
                self.assertEqual(result.status_code,200)
                manager.worker.result(timeout=10)
                self.assertEqual(manager.translation["phase"],"completed")
                self.assertNotIn(SECRET,client.get("/api/status").text)
            finally:
                client.close()

class TranslationWorkflowTests(unittest.TestCase):
    def test_dedup_cache_persistence_and_original_comparison_export(self):
        with MockProvider() as provider,tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"results.json"
            manager=TaskManager(path)
            tasks=[TASK,{**TASK,"id":"second","country":"uk"}]
            manager.start_run(tasks)
            raw=[{"review_id":"R1","title":"Works","body":"Good\nproduct","rating":5},{"review_id":"R2","title":"","body":"Another review"}]
            for task in tasks:
                manager.add_result(task["id"],copy.deepcopy(raw))
            manager.mark_complete()
            settings=provider.settings()
            manager.start_translation(settings);translate_results(manager,settings)
            self.assertEqual(manager.translation["succeeded"],4)
            self.assertEqual(len(provider.calls),1)
            self.assertEqual(len(json.loads(provider.calls[0]["body"]["messages"][1]["content"])["reviews"]),2)
            for values in manager.results.values():
                self.assertEqual(values[0]["body"],raw[0]["body"])
                self.assertEqual(values[0]["translation"]["body"],"Translated Good\nproduct")
            self.assertNotIn(SECRET,path.read_text(encoding="utf-8"))
            restored=TaskManager(path)
            self.assertEqual(restored.translation["phase"],"completed")
            restored.start_translation(settings);translate_results(restored,settings)
            self.assertEqual(len(provider.calls),1)
            changed=provider.settings(target_language="ja")
            restored.start_translation(changed);translate_results(restored,changed)
            self.assertEqual(len(provider.calls),2)
            workbook=load_workbook(io.BytesIO(server.build_excel(restored.tasks,restored.results,restored.outcomes)))
            rows=list(workbook[TASK["asin"]+"_us"].values)
            record=dict(zip(rows[0],rows[1]))
            self.assertEqual(record["正文"],"Good\nproduct")
            self.assertEqual(record["译文正文"],"Translated Good\nproduct")
            self.assertEqual(record["翻译目标语言"],"ja")
            self.assertNotIn(SECRET,str(record))

    def test_failure_then_retry_only_failed_without_fake_translation(self):
        with MockProvider() as provider:
            provider.mode="partial"
            manager=TaskManager();seed(manager)
            manager.results["first"].append({"review_id":"R2","title":"","body":"bad"})
            settings=provider.settings(batch_size=1)
            manager.start_translation(settings);translate_results(manager,settings)
            self.assertEqual(manager.translation["phase"],"partial")
            self.assertEqual(manager.translation["succeeded"],1)
            self.assertEqual(manager.results["first"][1]["translation"]["body"],"")
            provider.mode="normal"
            count=len(provider.calls)
            manager.start_translation(settings);translate_results(manager,settings)
            self.assertEqual(len(provider.calls),count+1)
            self.assertEqual(manager.translation["phase"],"completed")

    def test_fatal_auth_makes_one_call_and_keeps_all_originals(self):
        with MockProvider() as provider:
            provider.mode="unauthorized"
            manager=TaskManager();seed(manager)
            manager.results["first"] += [{"review_id":str(i),"title":"","body":str(i)} for i in range(10)]
            manager.start_translation(provider.settings(batch_size=1));translate_results(manager,provider.settings(batch_size=1))
            self.assertEqual(len(provider.calls),1)
            self.assertEqual(manager.translation["failed"],11)
            self.assertEqual(manager.translation["phase"],"failed")
            self.assertNotIn(SECRET,json.dumps(manager.snapshot()))
            self.assertEqual(manager.results["first"][0]["body"],"good")

    def test_cancel_mutual_exclusion_and_interrupted_restore(self):
        with tempfile.TemporaryDirectory() as folder,MockProvider() as provider:
            manager=TaskManager(Path(folder)/"results.json");seed(manager)
            manager.start_translation(provider.settings())
            with self.assertRaises(ValueError):
                manager.start_run([TASK])
            with self.assertRaises(ValueError):
                manager.start_login("us")
            manager.apply_translations({},0,0,0)
            restored=TaskManager(manager.persistence_path)
            self.assertFalse(restored.running)
            self.assertEqual(restored.translation["phase"],"interrupted")
            self.assertEqual(restored.phase,"completed")
            manager.stop();translate_results(manager,provider.settings())
            self.assertEqual(manager.translation["phase"],"cancelled")
            self.assertEqual(provider.calls,[])
            self.assertEqual(manager.results["first"][0]["body"],"good")

    def test_cancel_inflight_does_not_apply_late_response(self):
        with MockProvider() as provider:
            manager=TaskManager();seed(manager)
            manager.start_translation(provider.settings())
            def response(settings,items):
                manager.stop()
                return {"0":{"title":"late","body":"late"}}
            with patch.object(translation,"_request",side_effect=response):
                translate_results(manager,provider.settings())
            self.assertEqual(manager.translation["done"],0)
            self.assertNotIn("translation",manager.results["first"][0])
            self.assertEqual(manager.translation["phase"],"cancelled")

    def test_automatic_translation_starts_after_browser_close(self):
        with MockProvider() as provider:
            manager=TaskManager();manager.start_run([TASK])
            closed=[]
            request=translation._request
            def check(settings,items):
                self.assertEqual(closed,[True])
                self.assertEqual(manager.phase,"translating")
                return request(settings,items)
            with patch.object(server,"crawl",return_value=[{"title":"Source","body":"Source body"}]),patch.object(server.browser_mod,"close",side_effect=lambda:closed.append(True)),patch.object(translation,"_request",side_effect=check):
                server.BROWSER_WORKER.submit(server._worker_run,[TASK],manager,provider.settings()).result(timeout=10)
            self.assertEqual(manager.translation["phase"],"completed")
            self.assertEqual(manager.phase,"completed")
            self.assertFalse(manager.running)

    def test_export_formula_safety_applies_to_translations(self):
        manager=TaskManager();seed(manager)
        manager.results["first"][0]["translation"]={"title":"=SUM(1,2)","body":"+unsafe","status":"success"}
        workbook=load_workbook(io.BytesIO(server.build_excel(manager.tasks,manager.results)))
        rows=list(workbook[TASK["asin"]].values)
        row=dict(zip(rows[0],rows[1]))
        self.assertTrue(row["译文标题"].startswith("'="))
        self.assertTrue(row["译文正文"].startswith("'+"))


class TranslationAPITests(unittest.TestCase):
    def setUp(self):
        self.manager=TaskManager()
        self.patch=patch.object(server,"manager",self.manager);self.patch.start()
        self.client=TestClient(server.app)

    def tearDown(self):
        self.client.close();self.patch.stop()

    def test_connection_and_existing_results_api(self):
        with MockProvider() as provider:
            self.assertEqual(self.client.post("/api/translate",json=provider.body()).status_code,409)
            result=self.client.post("/api/llm/test",json=provider.body())
            self.assertEqual(result.status_code,200)
            self.assertNotIn(SECRET,result.text)
            seed(self.manager)
            result=self.client.post("/api/translate",json=provider.body())
            self.assertEqual(result.status_code,200)
            self.manager.worker.result(timeout=10)
            snapshot=self.client.get("/api/status")
            self.assertEqual(snapshot.json()["translation"]["phase"],"completed")
            self.assertNotIn(SECRET,snapshot.text)

    def test_start_request_auto_translation(self):
        with MockProvider() as provider,patch.object(server,"crawl",return_value=[{"title":"Title","body":"Body"}]):
            result=self.client.post("/api/start",json={"asins":[TASK["asin"]],"translation":provider.body()})
            self.assertEqual(result.status_code,200)
            self.manager.worker.result(timeout=10)
            self.assertEqual(self.manager.translation["succeeded"],1)
            self.assertNotIn(SECRET,json.dumps(self.manager.tasks))

    def test_provider_and_validation_errors_do_not_expose_keys(self):
        with MockProvider() as provider:
            provider.mode="unauthorized"
            response=self.client.post("/api/llm/test",json=provider.body())
            self.assertEqual(response.status_code,400)
            self.assertNotIn(SECRET,response.text)
            response=self.client.post("/api/start",json={"asins":[TASK["asin"]],"translation":provider.body(base_url=provider.base+"?api_key="+SECRET)})
            self.assertEqual(response.status_code,422)
            self.assertNotIn(SECRET,response.text)


class TranslationUITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pw=sync_playwright().start();cls.engine=cls.pw.chromium.launch(headless=True)

    @classmethod
    def tearDownClass(cls):
        cls.engine.close();cls.pw.stop()

    def setUp(self):
        self.context=self.engine.new_context()
        self.page=self.context.new_page()
        self.page.set_default_timeout(3000)

    def tearDown(self):
        self.context.close()

    load_frontend=old_tests.BrowserTests.load_frontend

    def test_anthropic_form_protocol_switch_draft_and_payload(self):
        self.load_frontend(origin=True)
        self.page.locator('#enableTranslation').check()
        self.page.locator('#llmKey').fill(SECRET)
        self.page.locator('#llmProvider').select_option('anthropic')
        self.assertEqual(self.page.locator('#llmKey').input_value(),'')
        self.assertEqual(self.page.locator('#llmUrl').input_value(),'https://api.anthropic.com')
        self.assertTrue(self.page.locator('#llmMaxTokens').is_visible())
        self.page.locator('#llmModel').fill('fixture-anthropic-model')
        self.page.locator('#llmMaxTokens').fill('8192')
        self.page.locator('#llmKey').fill(SECRET)
        self.page.wait_for_function("JSON.parse(localStorage.getItem('reviewStudio.settings')||'{}').translation?.max_tokens==='8192'")
        self.assertNotIn(SECRET,self.page.evaluate("localStorage.getItem('reviewStudio.settings')"))
        self.page.locator('#languageSelect').select_option('en')
        self.assertIn('Anthropic',self.page.locator('#llmProtocolBadge').inner_text())
        self.assertEqual(self.page.locator('#llmKey').input_value(),SECRET)
        settings=self.page.evaluate('translationSettings()')
        self.assertEqual(settings['provider'],'anthropic')
        self.assertEqual(settings['max_tokens'],8192)
        self.page.reload()
        self.page.evaluate("restoreTranslationDraft(JSON.parse(localStorage.getItem('reviewStudio.settings')).translation)")
        self.assertEqual(self.page.locator('#llmProvider').input_value(),'anthropic')
        self.assertEqual(self.page.locator('#llmMaxTokens').input_value(),'8192')
        self.assertEqual(self.page.locator('#llmKey').input_value(),'')
        self.page.locator('#llmProvider').select_option('openai')
        self.assertTrue(self.page.locator('#llmAnthropicField').is_hidden())
        self.assertNotIn('max_tokens',self.page.evaluate('translationSettings()'))

    def test_llm_form_bilingual_draft_never_persists_key(self):
        self.load_frontend(origin=True)
        self.assertFalse(self.page.locator("#enableTranslation").is_checked())
        self.page.locator("#enableTranslation").check()
        self.page.locator("#llmUrl").fill("http://127.0.0.1:12345/v1")
        self.page.locator("#llmModel").fill("fixture")
        self.page.locator("#llmKey").fill(SECRET)
        self.page.locator("#llmLanguage").select_option("ja")
        self.assertEqual(self.page.locator("#llmKey").get_attribute("type"),"password")
        self.page.wait_for_function("JSON.parse(localStorage.getItem('reviewStudio.settings')||'{}').translation?.model==='fixture'")
        draft=self.page.evaluate("localStorage.getItem('reviewStudio.settings')")
        self.assertNotIn(SECRET,draft)
        self.page.locator("#languageSelect").select_option("en")
        self.assertEqual(self.page.locator("#llmLanguage").input_value(),"ja")
        self.assertEqual(self.page.locator("#llmKey").input_value(),SECRET)
        self.page.reload()
        self.page.evaluate("restoreTranslationDraft(JSON.parse(localStorage.getItem('reviewStudio.settings')).translation)")
        self.assertEqual(self.page.locator("#llmKey").input_value(),"")
        self.assertEqual(self.page.locator("#llmModel").input_value(),"fixture")
        for width in (1440,800,420,360):
            self.page.set_viewport_size({"width":width,"height":950})
            self.assertFalse(self.page.evaluate("document.documentElement.scrollWidth>innerWidth"))

    def test_original_translation_comparison_search_and_expand(self):
        self.load_frontend()
        manager=TaskManager();seed(manager,body="original text "*30)
        manager.results["first"][0]["translation"]={"title":"<script>unsafe</script>","body":"译文针 "*120,"target_language":"zh-CN","status":"success"}
        self.page.evaluate("s=>applySnapshot(s)",manager.snapshot())
        self.assertIn("original text",self.page.locator(".review-text").inner_text())
        self.assertIn("译文针",self.page.locator(".review-translation-text").inner_text())
        self.assertEqual(self.page.locator(".review-translated script").count(),0)
        self.page.locator(".expand-review").click()
        self.assertTrue(self.page.locator(".review-text").evaluate("e=>e.classList.contains('expanded')"))
        self.assertTrue(self.page.locator(".review-translation-text").evaluate("e=>e.classList.contains('expanded')"))
        self.page.locator("#reviewSearch").fill("译文针")
        self.assertEqual(self.page.locator(".review-row").count(),1)
        self.page.locator("#languageSelect").select_option("en")
        self.assertIn("Translation",self.page.locator(".review-translated .comparison-label").inner_text())
        self.assertIn("译文针",self.page.locator(".review-translation-text").inner_text())

    def test_auto_payload_and_translation_progress(self):
        self.load_frontend()
        self.page.locator("#asins").fill(TASK["asin"])
        self.page.locator("#enableTranslation").check()
        self.page.locator("#llmUrl").fill("http://127.0.0.1:12345/v1")
        self.page.locator("#llmModel").fill("fixture")
        self.page.locator("#llmKey").fill(SECRET)
        self.page.evaluate("""()=>{window.fetch=async(url,options)=>{
          if(url==='/api/start')window.sentBody=JSON.parse(options.body);
          return {ok:true,json:async()=>url==='/api/start'?{ok:true}:{...state,instance_id:'fixture',revision:99}};
        }}""")
        self.page.locator("#startBtn").click()
        self.page.wait_for_function("window.sentBody!==undefined")
        self.assertEqual(self.page.evaluate("window.sentBody.translation.api_key"),SECRET)
        self.page.evaluate("""()=>applySnapshot({...state,instance_id:'translation',revision:1,running:true,phase:'translating',
          translation:{phase:'running',done:2,total:10,succeeded:2,failed:0,target_language:'zh-CN',model:'fixture'}})""")
        self.assertIn("正在翻译",self.page.locator("#translationTitle").inner_text())
        self.assertTrue(self.page.locator("#llmKey").is_disabled())
        self.assertTrue(self.page.locator("#stopTranslationBtn").is_visible())
